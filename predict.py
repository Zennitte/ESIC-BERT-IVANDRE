"""Fill an unlabeled XLSX using the included ensemble; preserve original input."""
import argparse
from collections import Counter
import json
import os
from pathlib import Path
import time
from datetime import datetime,timezone
import numpy as np
import openpyxl
import torch
from model import ROOT,Ensemble,digest

def atomic_json(path,value):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_suffix(path.suffix+'.tmp');tmp.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');os.replace(tmp,path)

def predict(input_path,output_path,sheet_name='test1',batch_size=4,device='cpu',artifacts=ROOT/'artifacts',checkpoint_rows=64):
    started=time.perf_counter();input_path=Path(input_path);output_path=Path(output_path)
    if input_path.resolve()==output_path.resolve():raise ValueError('Input and output must differ')
    input_hash=digest(input_path);artifact_hash=digest(Path(artifacts)/'manifest.json')
    workbook=openpyxl.load_workbook(input_path)
    sheet=workbook[sheet_name] if sheet_name else workbook.active
    headers={str(cell.value).strip():cell.column for cell in sheet[1] if cell.value is not None}
    if not {'resp_text','clarity'}<=set(headers):raise ValueError('Expected resp_text and clarity columns')
    tc,lc=headers['resp_text'],headers['clarity']
    rows=[i for i in range(2,sheet.max_row+1) if sheet.cell(i,tc).value is not None]
    if not rows:raise ValueError('No texts found')
    if any(sheet.cell(i,lc).value not in (None,'') for i in rows):raise ValueError('Input clarity must be empty')
    output_path.parent.mkdir(parents=True,exist_ok=True)
    state_path=output_path.with_suffix('.checkpoint.json')
    state=dict(input_sha256=input_hash,artifacts_manifest_sha256=artifact_hash,sheet=sheet.title,rows=rows,predictions=[],probabilities=[],status='running')
    if state_path.exists():
        previous=json.loads(state_path.read_text(encoding='utf-8'))
        for key in ['input_sha256','artifacts_manifest_sha256','sheet','rows']:
            if previous[key]!=state[key]:raise ValueError(f'Resume input changed: {key}')
        state=previous
    completed=len(state['predictions'])
    if completed>len(rows) or completed!=len(state['probabilities']):raise ValueError('Invalid saved checkpoint')
    classes=json.loads((Path(artifacts)/'pipeline.json').read_text(encoding='utf-8'))['class_order']
    if completed<len(rows):
        model=Ensemble(artifacts,device)
        for start in range(completed,len(rows),checkpoint_rows):
            current=rows[start:start+checkpoint_rows]
            texts=[str(sheet.cell(i,tc).value) for i in current]
            p=model.predict_proba(texts,batch_size)
            state['predictions'].extend(classes[i] for i in p.argmax(1));state['probabilities'].extend(p.tolist())
            state.update(completed_rows=len(state['predictions']),total_rows=len(rows),updated_utc=datetime.now(timezone.utc).isoformat())
            atomic_json(state_path,state)
            print(f'PREDICTION {len(state["predictions"])}/{len(rows)} checkpoint_saved',flush=True)
            if (ROOT/'PAUSE_REQUESTED').exists():
                print('PAUSED; remove PAUSE_REQUESTED and rerun the same command',flush=True);return None
        del model
        if device=='cuda':torch.cuda.empty_cache()
    probabilities=np.asarray(state['probabilities']);predicted=state['predictions']
    if probabilities.shape!=(len(rows),3) or not np.isfinite(probabilities).all() or not np.allclose(probabilities.sum(1),1,atol=1e-5):raise ValueError('Saved probability check failed')
    if predicted!=[classes[i] for i in probabilities.argmax(1)]:raise ValueError('Saved labels differ from argmax')
    for i,label in zip(rows,predicted,strict=True):sheet.cell(i,lc).value=label
    tmp=output_path.with_name(output_path.stem+'.tmp.xlsx');workbook.save(tmp);workbook.close();os.replace(tmp,output_path)
    # Verify every worksheet cell: only clarity for predicted rows may change.
    original=openpyxl.load_workbook(input_path,read_only=True,data_only=False)
    restored=openpyxl.load_workbook(output_path,read_only=True,data_only=False)
    if original.sheetnames!=restored.sheetnames:raise ValueError('Worksheet names changed')
    row_to_label=dict(zip(rows,predicted,strict=True));verified=0
    for name in original.sheetnames:
        a,b=original[name],restored[name]
        if (a.max_row,a.max_column)!=(b.max_row,b.max_column):raise ValueError('Worksheet dimensions changed')
        for index,(before,after) in enumerate(zip(a.iter_rows(values_only=True),b.iter_rows(values_only=True),strict=True),1):
            expected=list(before)
            if name==sheet.title and index in row_to_label:expected[lc-1]=row_to_label[index]
            if tuple(expected)!=tuple(after):raise ValueError(f'Unexpected cell change at {name}/{index}')
            verified+=len(before)
    original.close();restored.close()
    if digest(input_path)!=input_hash:raise ValueError('Original input changed')
    state['status']='complete';atomic_json(state_path,state)
    metadata=dict(status='complete',rows=len(rows),input=str(input_path),output=str(output_path),sheet=sheet.title,
        input_sha256=input_hash,output_sha256=digest(output_path),artifacts_manifest_sha256=artifact_hash,
        class_distribution=dict(Counter(predicted)),device=device,batch_size=batch_size,
        class_order=classes,weights=[0.5,0.5],new_training=False,labels_provided=False,accuracy=None,
        all_cells_verified=verified,input_preserved=True,probabilities_verified=True,
        wall_seconds=time.perf_counter()-started,completed_utc=datetime.now(timezone.utc).isoformat())
    atomic_json(output_path.with_suffix('.metadata.json'),metadata)
    print(json.dumps(metadata,ensure_ascii=False),flush=True)
    return metadata

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input',type=Path,default=ROOT/'data/test.xlsx')
    parser.add_argument('--output',type=Path,default=ROOT/'predictions/test_predito.xlsx')
    parser.add_argument('--sheet',default='test1');parser.add_argument('--batch-size',type=int,default=4)
    parser.add_argument('--device',choices=['cpu','cuda'],default='cpu')
    parser.add_argument('--artifacts',type=Path,default=ROOT/'artifacts')
    args=parser.parse_args()
    if args.batch_size<1:parser.error('--batch-size must be positive')
    result=predict(args.input,args.output,args.sheet,args.batch_size,args.device,args.artifacts)
    if result is None:raise SystemExit(75)

if __name__=='__main__':main()
